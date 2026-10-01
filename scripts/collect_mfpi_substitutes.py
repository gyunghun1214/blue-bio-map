"""verified-pilot-3.3 MFPI substitutes: FAO/INFOODS uFiSh1.0 items that match one of the 30 species at species,
genus or family level (diagram stage 3 MFPI, "섭취위치 / 유사종 대체치 구분").

Reads the workbook downloaded from UFISH_URL (sha256 checked), the 30 species (dist/assessments.json), their WoRMS
genus and family (research/verified-indices/taxonomy.json) and WoRMS synonyms (the MBPI link snapshot). Stores only
the matched raw items: ID, name, taxon label, consumed part, protein/iron/zinc per 100 g edible portion, the
component documentation code and n. Which item fills which gap is decided by the builder (config nutrition.substitutes).

    PYTHONUTF8=1 python scripts/collect_mfpi_substitutes.py --ufish <uFiSh1.0.xlsx>
"""
import argparse
import hashlib
import json
import re
from datetime import date
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
UFISH_URL = "https://www.fao.org/3/I6655EN/uFiSh1.0.xlsx"
UFISH_SHA256 = "66ce0037ce1f9296102d190e4aafc08aa30d658fd843a7591a9abdcdc787be12"  # downloaded 2026-09-30
COMPONENTS = {"protein_g": "PROTCNT(g)", "iron_mg": "FE(mg)", "zinc_mg": "ZN(mg)"}
PARTS = ("fillet", "flesh", "muscle", "whole")  # the consumed part as the item name states it


def overview(ws) -> dict[str, dict]:
    """Food item ID -> taxon from '02 Overview Species' (raw-food column only). '∆' marks a pooled genus or family entry."""
    items = {}
    for r in ws.iter_rows(min_row=5, values_only=True):
        if not (r[1] and r[4] and r[5]) or not re.fullmatch(r"\d{6}(-\d{6})?", str(r[5]).strip()):
            continue
        family, name = str(r[2]).strip(), " ".join(str(r[4]).split())
        if "∆" in family:
            level = "above_family" if "," in family else "family"
            taxon = family.replace("∆", "").strip()
        elif "spp∆" in name:
            level, taxon = "genus", name.split()[0]
        else:
            level, taxon = "species", re.match(r"[A-Z][a-z]+ [a-z]+", name).group(0)
        first, last = (int(x) for x in (str(r[5]).split("-") * 2)[:2])
        for i in range(first, last + 1):
            items[f"{i:06d}"] = {"level": level, "taxon": taxon, "family": family.replace("∆", "").strip(), "label": name}
    return items


def stat_docs(ws) -> dict[str, dict]:
    """Food item ID -> {column: (doc code, n)} from '05 NV_stat', plus the edible-factor comment."""
    header, out, current = None, {}, None
    for r in ws.iter_rows(values_only=True):
        if header is None:
            header = {h: i for i, h in enumerate(r) if h}
            continue
        key = str(r[0]).strip() if r[0] is not None else ""
        if re.fullmatch(r"\d{6}", key):
            current = out.setdefault(key, {"doc": {}, "n": {}, "comment": None})
        elif current is not None and key == "Documentation at component level":
            current["doc"] = {c: (str(r[i]).strip() if r[i] else None) for c, i in header.items()}
        elif current is not None and key == "n":
            current["n"] = {c: r[i] for c, i in header.items()}
        elif current is not None and key == "Comment":
            current["comment"] = r[header["EDIBLE"]]
    return out


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--ufish", type=Path, required=True)
    parser.add_argument("--out", type=Path)
    args = parser.parse_args()
    raw = args.ufish.read_bytes()
    if hashlib.sha256(raw).hexdigest() != UFISH_SHA256:
        raise SystemExit("uFiSh workbook differs from the recorded download (sha256)")
    report = json.loads((ROOT / "dist/assessments.json").read_text(encoding="utf-8"))
    taxonomy = json.loads((ROOT / "research/verified-indices/taxonomy.json").read_text(encoding="utf-8"))["species"]
    links = json.loads((ROOT / "research/verified-indices/snapshots/mbpi-links-2026-09-30.json").read_text(encoding="utf-8"))
    synonyms = {s["aphia_id"]: set(s["worms_synonyms"]) for s in links["species"]}
    targets = []
    for s in report["species"] + report["candidate_species"]:
        t = taxonomy[str(s["aphia_id"])]
        names = {s["scientific_name"], *synonyms.get(s["aphia_id"], ())}  # whole names: a trinomial synonym never makes its parent species a match
        targets.append({"aphia_id": s["aphia_id"], "names": names,
                        "genus": t["genus"], "family": t["family"]})

    wb = openpyxl.load_workbook(args.ufish, read_only=True, data_only=True)
    taxa, docs = overview(wb["02 Overview Species"]), stat_docs(wb["05 NV_stat (per 100 g EP)"])
    rows = list(wb["04 NV_sum (per 100 g EP)"].iter_rows(values_only=True))
    col = {h: i for i, h in enumerate(rows[0]) if h}
    items = []
    for r in rows[2:]:
        fid = str(r[0]).strip() if r[0] is not None else ""
        if fid not in taxa or r[col["State of food"]] != "r":
            continue
        taxon = taxa[fid]
        genus = taxon["taxon"].split()[0] if taxon["level"] in ("species", "genus") else None
        matched = {}
        for t in targets:  # the closest level only; an entry above family level is never a substitute
            if taxon["level"] == "species" and taxon["taxon"] in t["names"]:
                matched[t["aphia_id"]] = "species"
            elif genus and genus == t["genus"]:
                matched[t["aphia_id"]] = "genus"
            elif taxon["level"] != "above_family" and taxon["family"] == t["family"]:
                matched[t["aphia_id"]] = "family"
        if not matched:
            continue
        name = str(r[col["Food name in English"]]).strip()
        doc = docs.get(fid, {"doc": {}, "n": {}, "comment": None})
        items.append({
            "food_item_id": fid, "food_name": name, "taxon_label": taxon["label"], "taxon_level": taxon["level"],
            "family": taxon["family"], "habitat": r[col["Habitat"]], "state": "raw",
            "part": next((p for p in PARTS if re.search(rf"\b{p}\b", name.lower())), None),
            "edible_factor_comment": doc["comment"], "matches": {str(k): v for k, v in sorted(matched.items())},
            "components": {k: {"value": r[col[c]] if isinstance(r[col[c]], (int, float)) else None,
                               "doc": doc["doc"].get(c), "n": doc["n"].get(c) if isinstance(doc["n"].get(c), int) else None}
                           for k, c in COMPONENTS.items()}})
    today = date.today().isoformat()
    snapshot = {
        "snapshot_date": today,
        "sources": {"ufish_1_workbook": {
            "provider": "FAO/INFOODS", "title": "Global Food Composition Database for Fish and Shellfish (uFiSh1.0), 02 Overview Species, 04 NV_sum and 05 NV_stat",
            "version": f"uFiSh1.0 (December 2016); workbook sha256 {UFISH_SHA256}", "url": UFISH_URL, "accessed": today,
            "license": "FAO copyright; non-commercial research and education use with attribution",
            "terms": "© FAO 2016 (user guide i6655en.pdf p.4, read 2026-10-01): material may be copied, downloaded and printed for private study, research and teaching purposes, or for use in non-commercial products or services, provided that appropriate acknowledgement of FAO as the source and copyright holder is given and that FAO's endorsement is not implied; translation, adaptation, resale and other commercial use rights must be requested from FAO. Cite as: FAO (2016). FAO/INFOODS Global Food Composition Database for Fish and Shellfish Version 1.0 - uFiSh1.0. Rome. Only matched item IDs, three values, documentation codes and n are stored.",
            "user_guide": "https://www.fao.org/3/I6655EN/i6655en.pdf",
            "documentation_codes": {"a": "analytical value", "r": "value taken from a reference food composition dataset",
                                    "ar": "mix of analytical and reference-dataset values", "c": "calculated in the datasheet (e.g. protein from nitrogen)",
                                    "e": "estimated or borrowed from a similar food"}}},
        "basis": "per 100 g edible portion, raw", "components": COMPONENTS,
        "match_rule": "species: the uFiSh scientific name equals the WoRMS accepted name or a WoRMS synonym; genus: the first word of the uFiSh scientific name as published equals the WoRMS genus of the species (not resolved to the accepted combination, so uFiSh 'Crassostrea gigas' does not match genus Magallana); family: the uFiSh FAMILY column equals the WoRMS family; pooled entries spanning several families (e.g. 093034 'Sepiidae, Sepiolidae') are never matched",
        "items": items}
    out = args.out or ROOT / f"research/verified-indices/snapshots/ufish-substitutes-{today}.json"
    out.write_text(json.dumps(snapshot, ensure_ascii=False, indent=1, sort_keys=True) + "\n", encoding="utf-8")
    print(f"Wrote {out}: {len(items)} matched raw items")


if __name__ == "__main__":
    main()
